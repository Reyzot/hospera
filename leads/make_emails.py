"""Genera secuencia de 3 emails personalizados por hotel (datos reales de Google Maps)."""
import csv, statistics, re, collections, openpyxl
from openpyxl.styles import Alignment, Font
rows = list(csv.DictReader(open("florida_hotels_independientes.csv")))
num = lambda v: int(float(v)) if v not in ("", None) else 0
flt = lambda v: float(v) if v not in ("", None) else 0.0
bycity = collections.defaultdict(list)
for r in rows: bycity[r["city"]].append(r)
DEMO = "https://hosperai.es/demo"

def short(name):
    n = re.sub(r"\s*[\(\|–-].*$", "", name).strip()
    return n if len(n) <= 40 else n[:40].rsplit(" ", 1)[0]

NAMES = set("""ana andrew arthur bill chastity heather leila lynn marilyn monica tracey victoria john mike michael david chris
james robert mary patricia jennifer linda susan karen lisa nancy betty sandra donna carol sharon michelle laura sarah kim jessica
amy angela melissa rebecca stephanie nicole elizabeth julie joyce anna emily kelly christine debbie deb cathy kathy tom tim steve
paul mark brian kevin jason jeff gary eric scott greg frank ray dan dave ron joe jim bob rick richard peter george ed carlos maria
jose luis juan jorge ana sofia daniel alex sam ben matt nick tony pat terry chuck doug roger wendy tina diane janet pam jan ann
anne jane jill gail beth barbara helen ruth kate katie jen jenny amanda ashley brittany megan rachel andrea sue bonnie brenda""".split())
def first_name(email, hotel=""):
    local = email.split("@")[0].lower()
    joined = re.sub(r"[^a-z]", "", hotel.lower())
    if local in joined or any(w in local for w in ("inn","hotel","beach","house","resort","cottage","suite","villa","lodge","motel","sea","bay","key","island","palm","ocean","sun","stay","club")):
        return None
    if re.fullmatch(r"[a-z]{3,12}", local) and local not in {"info","reservations","reservation","stay","hello","contact","frontdesk","office","sales","manager","booking","bookings","reserve","admin","events","weddings","concierge","guest","guests","welcome","rooms","inquiries","mail","email","team","general","relax","destinations","rentals","resort","hotel","inn","gm"}:
        return local.capitalize() if local in NAMES else None
    return None

out = []
for r in rows:
    if re.search(r"(?i)private |vacation|rentals", r["name"] + r["email"]): continue
    if not r["email"] or re.match(r"(janesmith|johndoe|impallari|tennisshopcc|example|test|name)@", r["email"]): continue
    rev, rat = num(r["reviews"]), flt(r["rating"])
    peers = sorted([p for p in bycity[r["city"]] if p["name"] != r["name"] and not re.search(r"(?i)luxury collection|marriott|tribute|autograph|vacation|private|rental", p["name"])], key=lambda p: -num(p["reviews"]))
    # comparar con un negocio parecido: mismo tipo y como mucho 8 veces sus reseñas (no un resort gigante vs un B&B)
    similar = [p for p in peers if p["type"] == r["type"] and num(p["reviews"]) <= max(num(r["reviews"]) * 8, 150)]
    leader = similar[0] if similar else None
    med = int(statistics.median([num(p["reviews"]) for p in bycity[r["city"]]] or [0]))
    n = short(r["name"]); city = r["city"]; fn = first_name(r["email"], r["name"])
    hi = f"Hi {fn}," if fn else "Hi there,"
    lead_line = ""
    if leader and num(leader["reviews"]) > rev * 1.5:
        lead_line = f" {short(leader['name'])}, also in {city}, has {num(leader['reviews']):,}."
    if rev < med or (leader and num(leader["reviews"]) > rev * 2):
        seg = "A · pocas reseñas"
        subject = f"{rev:,} reviews"
        hook = f"{n} has {rev:,} Google reviews.{lead_line}"
    elif rat and rat < 4.5:
        seg = "B · nota mejorable"
        subject = f"{n} at {rat} stars"
        hook = f"{n} sits at {rat}★ on Google. Unhappy guests always write; happy ones rarely do."
    else:
        seg = "C · fuerte"
        subject = f"{n}'s reviews"
        hook = f"{n} has a great {rat}★ from {rev:,} Google reviews. Keeping that pace is the hard part."
    e1 = f"""{hi}

{hook}

Most happy guests just never get asked at the right moment. We fix that: after checkout they get a thank-you on WhatsApp with a one-tap Google link, and every new review comes with a reply already written for you.

I made a 38-second video of how it works. Want me to send it over?

Andreu
Hosperai"""
    e2 = f"""{hi}

In case it's easier than replying, here's the 38-sec video:
{DEMO}

No app, no dashboard, nothing extra for the front desk.

Andreu"""
    e3 = f"""{hi}

Should I close the loop on this? If reviews aren't a priority for {n} right now, just reply "not now" and I won't follow up.

Andreu"""
    out.append({"name": r["name"], "city": city, "email": r["email"], "segment": seg, "rating": rat, "reviews": rev,
                "city_median_reviews": med, "subject": subject, "email_1": e1, "subject_2": "Re: " + subject, "email_2": e2,
                "subject_3": "Re: " + subject, "email_3": e3, "phone": r["phone"], "website": r["website"]})

cols = list(out[0].keys())
with open("emails_personalizados.csv", "w", newline="") as f:
    w = csv.DictWriter(f, cols); w.writeheader(); w.writerows(out)
wb = openpyxl.load_workbook("florida_hotels_independientes.xlsx")
if "Emails" in wb.sheetnames: del wb["Emails"]
ws = wb.create_sheet("Emails", 0); ws.append(cols)
for c in ws[1]: c.font = Font(bold=True)
for o in out: ws.append([o[c] for c in cols])
for col, wd in zip("ABCDEFGHIJKLMNO", [30,16,30,18,7,8,9,34,70,30,50,30,50,16,34]): ws.column_dimensions[col].width = wd
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = "B2"; ws.auto_filter.ref = ws.dimensions
wb.save("florida_hotels_independientes.xlsx")
print(len(out), collections.Counter(o["segment"] for o in out))
for o in out[:1] + [x for x in out if x["segment"].startswith("B")][:1]:
    print("\n=====", o["subject"], "\n" + o["email_1"])
